"""Everything Miles can do on the PC. Each @tool is exposed to the AI as a callable function."""
import base64
import ctypes
import ctypes.wintypes as wt
import datetime as dt
import io
import json
import os
import re
import shutil
import socket
import subprocess
import threading
import time
import urllib.parse
import webbrowser
from pathlib import Path

import psutil
import requests

import config
from . import security, uia
from .apps import ALIASES, _norm
from .util import DATA, NO_WINDOW, folder_alias, known_folder, log, resolve_path, run_ps, user_search_roots

user32 = ctypes.windll.user32


class Ctx:
    apps = None                                   # AppIndex
    announce = staticmethod(lambda text: None)    # speak something unprompted (timers, scans)
    confirm = staticmethod(lambda q: False)       # ask the user yes/no by voice
    ask = staticmethod(lambda question, suggestion="": None)   # dialog box where the user types an answer
    status = staticmethod(lambda text: None)      # show progress on the HUD
    run_steps = staticmethod(lambda steps: ["Error: not available"] * len(steps))  # run plain-English commands
    side_call = None                              # (prompt, image_b64) -> answer, inside the current conversation
    cancelled = staticmethod(lambda: False)
    vision_model = None
    input_mode = "voice"                          # "voice" or "typed": how the current request arrived
    typed_answers: set = set()                    # what the user typed into ask boxes (trusted, never re-checked)
    request_text = ""                             # the user's current request, word for word
    snapshot = None                               # context.Snapshot: what was on screen when they asked
    last_write = None                             # {"hwnd", "steps", "t"} of the last in-place write (for undo)
    writer = None                                 # pending Writer: the brain's next message gets typed out
    plan: list = []
    active_brain = None


CTX = Ctx()
TOOLS: dict[str, dict] = {}


def tool(desc, props=None, required=(), confirm=None, enabled=None):
    def deco(fn):
        TOOLS[fn.__name__] = {
            "fn": fn, "confirm": confirm, "enabled": enabled,
            "schema": {"type": "function", "function": {
                "name": fn.__name__, "description": desc,
                "parameters": {"type": "object", "properties": props or {}, "required": list(required)}}}}
        return fn
    return deco


def S(desc, enum=None):
    d = {"type": "string", "description": desc}
    if enum:
        d["enum"] = enum
    return d


def N(desc):
    return {"type": "number", "description": desc}


def B(desc):
    return {"type": "boolean", "description": desc}


# Lite: the AI sees only these tools, with short descriptions. On a laptop CPU the model reads its prompt at
# ~50-150 tokens/s, so every word here costs start-up time and RAM. Everything else (volume, windows,
# media, power, protocols, security...) is handled instantly by the command router without the AI.
LITE_TOOLS = {
    "open_app": "Open any installed app or website by name.",
    "browser_tab": "Browser tabs: close/switch/mute/reload one tab by title, new, reopen, list. Never close the "
                   "whole browser for a tab.",
    "spotify": "Spotify: play a song/artist/playlist, pause, resume, next, previous, now_playing.",
    "open_website": "Open a URL.",
    "web_search": "Open a web search in the browser.",
    "research": "Look up current facts on the internet (news, prices, anything recent). Returns results.",
    "read_webpage": "Read a web page's text.",
    "write": "Write text (note, message, reply, poem, refined version of selected text). Then write the text "
             "itself as your next message. where: auto, selection (replace highlighted text), replace (whole box), "
             "notepad, word, clipboard.",
    "compose_email": "Write an email (you write subject and polished body) and open it in Gmail. send=true only "
                     "if asked to send.",
    "send_email": "Send the email open in the compose window.",
    "ask_user": "Ask the user to type a detail you don't know (email address, spelling). Never passwords.",
    "remember": "Save a lasting fact about the user (forget=true removes it).",
    "recall": "Search past conversations and saved facts.",
    "set_reminder": "Reminder/timer, or run a command later (at_time '18:00' or minutes).",
    "system_status": "PC health: CPU, RAM, disk, battery, top processes.",
    "file_action": "Files: open, list, read, write, copy, move/rename, delete, zip. Several paths: join with ' | '.",
    "find_files": "Find files or folders by name.",
    "type_text": "Type exact text into the focused box.",
    "press_keys": "Press keys, e.g. 'ctrl+s', 'enter', 'alt+tab'.",
    "read_screen": "List the buttons, links and fields in the active window.",
    "click_element": "Click or type into an element of the active window by its name.",
    "click_on_screen": "Click visible text anywhere on screen.",
    "look_at_screen": "Read the text visible on screen to answer questions about it.",
    "run_command": "Run a PowerShell command (anything else). Returns output.",
}


def _compact(schema: dict, desc: str) -> dict:
    """Same tool, shorter wording: the description above, parameter descriptions cut to a few words."""
    fn = schema["function"]
    props = {}
    for k, v in fn["parameters"]["properties"].items():
        p = {kk: vv for kk, vv in v.items() if kk != "description"}
        d = v.get("description", "")
        if d and "enum" not in v:
            p["description"] = d.split(";")[0].split("(")[0].split(" - ")[0].strip()[:50]
        props[k] = p
    return {"type": "function", "function": {"name": fn["name"], "description": desc,
                                             "parameters": {**fn["parameters"], "properties": props}}}


def schemas():
    return [_compact(TOOLS[n]["schema"], d) for n, d in LITE_TOOLS.items() if n in TOOLS]


def execute(name: str, args) -> str:
    t = TOOLS.get(name)
    if not t:
        return f"Error: there is no tool called {name}."
    if isinstance(args, str):
        try:
            args = json.loads(args or "{}")
        except json.JSONDecodeError:
            args = {}
    allowed = t["schema"]["function"]["parameters"]["properties"]
    args = {k: v for k, v in (args or {}).items() if k in allowed and v not in (None, "")}
    question = t["confirm"](args) if callable(t["confirm"]) else t["confirm"]
    if question and config.CONFIRM_DANGEROUS and not CTX.confirm(question):
        return "The user said no, so it was NOT done. Acknowledge briefly."
    if name in ("type_text", "press_keys", "mouse", "read_screen", "click_element", "manage_window",
                "look_at_screen", "take_screenshot", "write", "click_on_screen"):
        _await_launch()
    try:
        result = t["fn"](**args)
    except TypeError as e:
        return f"Error: wrong arguments ({e})."
    except Exception as e:
        log.exception("tool %s failed", name)
        return f"Error: {e}"
    return "Done." if result is None else str(result)


# =============================================================================
# Windows helpers
# =============================================================================
_WNDENUMPROC = ctypes.WINFUNCTYPE(ctypes.c_bool, wt.HWND, wt.LPARAM)
_SKIP_CLASSES = {"Progman", "WorkerW", "Shell_TrayWnd", "Shell_SecondaryTrayWnd", "TkTopLevel"}
_PROTECTED = {"explorer.exe", "svchost.exe", "csrss.exe", "winlogon.exe", "lsass.exe", "services.exe",
              "smss.exe", "wininit.exe", "dwm.exe", "system", "python.exe", "pythonw.exe", "ollama.exe",
              "ollama app.exe", "registry", "fontdrvhost.exe", "msmpeng.exe"}
EXE_HINTS = {
    "google chrome": "chrome.exe", "microsoft edge": "msedge.exe", "visual studio code": "code.exe",
    "file explorer": "explorer.exe", "word": "winword.exe", "excel": "excel.exe", "powerpoint": "powerpnt.exe",
    "outlook": "outlook.exe", "terminal": "windowsterminal.exe", "task manager": "taskmgr.exe",
    "calculator": "calculatorapp.exe", "settings": "systemsettings.exe", "notepad": "notepad.exe",
    "claude": "claude.exe", "whatsapp": "whatsapp.exe", "spotify": "spotify.exe", "discord": "discord.exe",
    "steam": "steam.exe", "obs studio": "obs64.exe", "vlc media player": "vlc.exe", "firefox": "firefox.exe",
    "brave": "brave.exe", "opera": "opera.exe", "paint": "mspaint.exe", "photos": "photos.exe",
}


def _cloaked(hwnd):
    v = ctypes.c_int(0)
    try:
        ctypes.windll.dwmapi.DwmGetWindowAttribute(hwnd, 14, ctypes.byref(v), ctypes.sizeof(v))
    except Exception:
        return False
    return v.value != 0


def windows():
    """Visible top-level windows -> list of (hwnd, title, process_name)."""
    out = []

    def cb(hwnd, _):
        if not user32.IsWindowVisible(hwnd) or user32.GetWindow(hwnd, 4):   # 4 = GW_OWNER
            return True
        n = user32.GetWindowTextLengthW(hwnd)
        if n == 0:
            return True
        cls = ctypes.create_unicode_buffer(64)
        user32.GetClassNameW(hwnd, cls, 64)
        if cls.value in _SKIP_CLASSES or _cloaked(hwnd):
            return True
        buf = ctypes.create_unicode_buffer(n + 1)
        user32.GetWindowTextW(hwnd, buf, n + 1)
        pid = wt.DWORD()
        user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
        try:
            pname = psutil.Process(pid.value).name().lower()
        except Exception:
            pname = "?"
        if buf.value in ("Program Manager", "Miles"):
            return True
        out.append((hwnd, buf.value, pname))
        return True
    user32.EnumWindows(_WNDENUMPROC(cb), 0)
    return out


_STOPWORDS = {"the", "a", "an", "it", "this", "that", "them", "these", "those", "all", "my", "one", "window",
              "app", "tab", "thing", "everything", "current", "please", "now"}


def _match_windows(query: str, allow_title: bool = True):
    """Windows belonging to an app (by process), or - only for specific enough words - by title.

    Returns (windows, how) where how is "process", "title" or None.
    """
    q = _norm(query)
    q = ALIASES.get(q, q)
    if not q or q in _STOPWORDS or len(q) < 3:
        return [], None
    exe = EXE_HINTS.get(q)
    q_compact = q.replace(" ", "")
    wins = windows()
    by_proc = [w for w in wins if (exe and w[2] == exe) or (len(q_compact) > 2 and q_compact in w[2].replace(".exe", ""))]
    if by_proc:
        return by_proc, "process"
    if allow_title and len(q) >= 4:
        by_title = [w for w in wins if q in w[1].lower()]
        if by_title:
            return by_title, "title"
    return [], None


def _is_browser(proc: str) -> bool:
    from .browser import BROWSERS
    return proc in BROWSERS


def focus_hwnd(hwnd):
    """Bring a window to the front. Windows blocks background apps from doing that, so: join the input queue of
    the window that has focus; if that isn't enough, press F24 (a key nothing uses). Not Alt - an Alt tap
    that lands in the target window opens its menu mode and swallows the next keys."""
    if user32.IsIconic(hwnd):
        user32.ShowWindow(hwnd, 9)
    fg = user32.GetForegroundWindow()
    if fg == hwnd:
        return
    fg_thread = user32.GetWindowThreadProcessId(fg, None)
    me = ctypes.windll.kernel32.GetCurrentThreadId()
    attached = bool(fg_thread) and fg_thread != me and user32.AttachThreadInput(me, fg_thread, True)
    user32.SetForegroundWindow(hwnd)
    user32.BringWindowToTop(hwnd)
    if attached:
        user32.AttachThreadInput(me, fg_thread, False)
    if user32.GetForegroundWindow() != hwnd:
        user32.keybd_event(0x87, 0, 0, 0)           # VK_F24
        user32.keybd_event(0x87, 0, 2, 0)
        user32.SetForegroundWindow(hwnd)
        user32.BringWindowToTop(hwnd)


def _foreground_title():
    hwnd = user32.GetForegroundWindow()
    buf = ctypes.create_unicode_buffer(512)
    user32.GetWindowTextW(hwnd, buf, 512)
    return hwnd, buf.value


# =============================================================================
# Apps & windows
# =============================================================================
SITES = {
    "youtube": "https://www.youtube.com", "gmail": "https://mail.google.com", "google": "https://www.google.com",
    "github": "https://github.com", "netflix": "https://www.netflix.com", "amazon": "https://www.amazon.in",
    "instagram": "https://www.instagram.com", "facebook": "https://www.facebook.com", "twitter": "https://x.com",
    "x": "https://x.com", "reddit": "https://www.reddit.com", "linkedin": "https://www.linkedin.com",
    "chatgpt": "https://chatgpt.com", "claude ai": "https://claude.ai", "google drive": "https://drive.google.com",
    "drive": "https://drive.google.com", "google maps": "https://maps.google.com", "maps": "https://maps.google.com",
    "wikipedia": "https://www.wikipedia.org", "stack overflow": "https://stackoverflow.com",
    "google docs": "https://docs.google.com", "google sheets": "https://sheets.google.com",
    "calendar": "https://calendar.google.com", "google calendar": "https://calendar.google.com",
    "youtube music": "https://music.youtube.com", "prime video": "https://www.primevideo.com",
    "hotstar": "https://www.hotstar.com", "flipkart": "https://www.flipkart.com", "twitch": "https://www.twitch.tv",
    "whatsapp web": "https://web.whatsapp.com", "outlook web": "https://outlook.live.com", "news": "https://news.google.com",
}


def site_url(name: str):
    n = _norm(name)
    n = re.sub(r"\s*(website|site|dot com|com)$", "", n).strip()
    if n in SITES:
        return SITES[n]
    if re.fullmatch(r"[\w-]+(\.[\w-]+)+(/\S*)?", name.strip().lower()):
        return "https://" + name.strip()
    return None


@tool("Open/launch any installed app, game or tool by name. Finds it automatically (Start menu, Store apps, "
      "desktop, CLI tools), e.g. 'chrome', 'claude code', 'spotify', 'vs code', 'steam', 'task manager'.",
      {"name": S("App name as spoken")}, ["name"])
def open_app(name: str, wait: bool = True):
    app, _ = CTX.apps.find(name)
    if not app:
        folder = folder_alias(name)
        if folder:
            os.startfile(folder)
            return f"Opened the {folder.name} folder."
        url = site_url(name)
        if url:
            webbrowser.open(url)
            return f"No app called {name}, so I opened {url} in the browser."
        sug = CTX.apps.suggestions(name)
        return (f"Couldn't find an app called '{name}'." +
                (f" Similar installed apps: {', '.join(sug)}." if sug else ""))
    before, _ = _foreground_title()
    CTX.apps.launch(app)
    _PENDING_LAUNCH[:] = [time.time(), before]       # next UI step waits for the window, not this call
    return f"Opened {app['name']}."


_PENDING_LAUNCH = []


def _await_launch():
    """If an app was just launched, wait (max ~4 s) for its window before typing/clicking into it."""
    if not _PENDING_LAUNCH:
        return
    started, before = _PENDING_LAUNCH
    _PENDING_LAUNCH.clear()
    end = started + 4.5
    while time.time() < end:
        if user32.GetForegroundWindow() != before:
            time.sleep(max(0.0, min(0.4, started + 0.8 - time.time())))   # let it finish drawing
            return
        time.sleep(0.1)


def _close_q(a):
    return f"Force-close {a.get('name')}? Unsaved work will be lost." if a.get("force") else None


@tool("Close an app (all its windows). force=true kills it without saving.",
      {"name": S("App or window name"), "force": B("Kill immediately")}, ["name"], confirm=_close_q)
def close_app(name: str, force: bool = False):
    q = _norm(name)
    if not q or q in _STOPWORDS or len(q) < 3:
        return f"'{name}' isn't an app name. Ask the user what to close."
    if re.search(r"\btabs?\b", q):                        # "youtube tab" is a browser tab, not an app
        from . import browser
        return browser.tab_action("close", re.sub(r"\btabs?\b", "", name).strip())
    wins, how = _match_windows(name)
    if how == "title" and all(_is_browser(w[2]) for w in wins):
        # the words match a web page, not an app: close just that tab, never the whole browser
        from . import browser
        return browser.tab_action("close", name)
    procs = {w[2] for w in wins}
    if not force and wins:
        for hwnd, _, _ in wins:
            user32.PostMessageW(hwnd, 0x0010, 0, 0)          # WM_CLOSE: polite close
        return f"Closed {len(wins)} window{'s' if len(wins) > 1 else ''} of {name}."
    if not procs:
        q = _norm(ALIASES.get(_norm(name), name)).replace(" ", "")
        exe = EXE_HINTS.get(ALIASES.get(_norm(name), _norm(name)))
        procs = {p.info["name"].lower() for p in psutil.process_iter(["name"])
                 if p.info["name"] and (p.info["name"].lower() == exe or q in p.info["name"].lower().replace(".exe", ""))}
    procs -= _PROTECTED
    if not procs:
        return f"{name} doesn't seem to be running."
    for p in procs:
        subprocess.run(["taskkill", "/IM", p, "/T"] + (["/F"] if force else []), capture_output=True,
                       creationflags=NO_WINDOW)
    return f"Closed {', '.join(procs)}."


@tool("Control a window: focus/switch to it, minimize, maximize, restore, close, snap_left, snap_right. "
      "Use name 'all' with minimize to show the desktop.",
      {"name": S("App or window title"), "action": S("What to do",
       ["focus", "minimize", "maximize", "restore", "close", "snap_left", "snap_right"])}, ["name", "action"])
def manage_window(name: str, action: str):
    import pyautogui
    if name.lower() in ("all", "everything", "all windows") and action == "minimize":
        pyautogui.hotkey("win", "d")
        return "Showing the desktop."
    how = "process"
    if name.lower() in ("this", "current", "active", "this window", "current window", "it"):
        hwnd, title = _foreground_title()
        pid = wt.DWORD()
        user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
        try:
            proc = psutil.Process(pid.value).name().lower()
        except Exception:
            proc = ""
        wins = [(hwnd, title, proc)]
    else:
        wins, how = _match_windows(name)
    if not wins:
        return f"No open window matches '{name}'."
    hwnd, title, proc = wins[0]
    if action == "close" and _is_browser(proc) and how == "title":
        from . import browser                     # a page title matched: close that tab, not the browser
        return browser.tab_action("close", name)
    if action == "focus":
        focus_hwnd(hwnd)
    elif action == "minimize":
        for w in wins:
            user32.ShowWindow(w[0], 6)
    elif action == "maximize":
        focus_hwnd(hwnd)
        user32.ShowWindow(hwnd, 3)
    elif action == "restore":
        user32.ShowWindow(hwnd, 9)
    elif action == "close":
        user32.PostMessageW(hwnd, 0x0010, 0, 0)
    elif action in ("snap_left", "snap_right"):
        focus_hwnd(hwnd)
        time.sleep(0.15)
        pyautogui.hotkey("win", "left" if action == "snap_left" else "right")
    return f"{action.replace('_', ' ').capitalize()}: {title}"


@tool("List open windows (title and process).")
def list_windows():
    ws = windows()
    if not ws:
        return "No windows open."
    fg = user32.GetForegroundWindow()
    return "\n".join(f"{'* ' if h == fg else ''}{t}  [{p}]" for h, t, p in ws[:40])


# =============================================================================
# Web & media
# =============================================================================
@tool("Open a website or URL in the browser (e.g. 'youtube', 'gmail', 'github.com/user').",
      {"url": S("Site name or URL")}, ["url"])
def open_website(url: str):
    target = site_url(url) or (url if url.startswith(("http://", "https://")) else "https://" + url.replace(" ", ""))
    webbrowser.open(target)
    return f"Opened {target}"


_SEARCH = {
    "google": "https://www.google.com/search?q={}", "youtube": "https://www.youtube.com/results?search_query={}",
    "bing": "https://www.bing.com/search?q={}", "wikipedia": "https://en.wikipedia.org/w/index.php?search={}",
    "amazon": "https://www.amazon.in/s?k={}", "github": "https://github.com/search?q={}",
    "maps": "https://www.google.com/maps/search/{}", "images": "https://www.google.com/search?tbm=isch&q={}",
    "news": "https://news.google.com/search?q={}", "flipkart": "https://www.flipkart.com/search?q={}",
    "stackoverflow": "https://stackoverflow.com/search?q={}", "spotify": "https://open.spotify.com/search/{}",
}


@tool("Search the web (opens results in the browser).",
      {"query": S("Search terms"),
       "site": S("Where to search. Default google; use another only if the user names it", list(_SEARCH))},
      ["query"])
def web_search(query: str, site: str = "google"):
    webbrowser.open(_SEARCH.get(site, _SEARCH["google"]).format(urllib.parse.quote_plus(query)))
    return f"Searching {site} for {query}."


@tool("Play a video on YouTube (top result). Only when the user asks for YouTube or a video; music goes to spotify.",
      {"query": S("What to play")}, ["query"])
def play_youtube(query: str):
    url = "https://www.youtube.com/results?search_query=" + urllib.parse.quote_plus(query)
    try:
        # stream the results page and stop at the first video id instead of downloading ~1 MB
        with requests.get(url, headers={"User-Agent": "Mozilla/5.0", "Accept-Language": "en"}, timeout=6,
                          stream=True) as r:
            r.encoding = "utf-8"
            buf = ""
            for chunk in r.iter_content(32768, decode_unicode=True):
                buf += chunk
                m = re.search(r'"videoId":"([\w-]{11})".{0,1500}?"title":\{"runs":\[\{"text":"(.*?)"\}', buf, re.S)
                if m:
                    webbrowser.open(f"https://www.youtube.com/watch?v={m.group(1)}")
                    return f"Playing {m.group(2)} on YouTube."
                if len(buf) > 1_500_000:
                    break
    except Exception as e:
        log.warning("youtube lookup failed: %s", e)
    webbrowser.open(url)
    return f"Opened YouTube results for {query}."


@tool("Media playback keys: play_pause, next, previous, stop.",
      {"action": S("Media action", ["play_pause", "next", "previous", "stop"])}, ["action"])
def media_control(action: str):
    import pyautogui
    pyautogui.press({"play_pause": "playpause", "next": "nexttrack", "previous": "prevtrack",
                     "stop": "stop"}.get(action, "playpause"))
    return f"Media: {action.replace('_', ' ')}."


def _volume_endpoint():
    try:
        import comtypes
        try:
            comtypes.CoInitialize()
        except OSError:
            pass
        from pycaw.pycaw import AudioUtilities, IAudioEndpointVolume
        dev = AudioUtilities.GetSpeakers()
        ep = getattr(dev, "EndpointVolume", None)
        if ep is None:                                      # older pycaw API
            iface = dev.Activate(IAudioEndpointVolume._iid_, comtypes.CLSCTX_ALL, None)
            ep = ctypes.cast(iface, ctypes.POINTER(IAudioEndpointVolume))
        return ep
    except Exception as e:
        log.debug("pycaw unavailable: %s", e)
        return None


@tool("System volume: set to a level, turn up/down, mute or unmute.",
      {"action": S("Volume action", ["set", "up", "down", "mute", "unmute"]),
       "level": N("0-100 for 'set', or step size for up/down (default 10)")}, ["action"])
def volume(action: str, level: float = None):
    import pyautogui
    ep = _volume_endpoint()
    if ep is not None:
        cur = round(ep.GetMasterVolumeLevelScalar() * 100)
        if action == "set" and level is not None:
            new = max(0, min(100, level))
        elif action == "up":
            new = min(100, cur + (level or 10))
        elif action == "down":
            new = max(0, cur - (level or 10))
        else:
            ep.SetMute(1 if action == "mute" else 0, None)
            return "Muted." if action == "mute" else f"Unmuted. Volume is {cur}%."
        ep.SetMute(0, None)
        ep.SetMasterVolumeLevelScalar(new / 100, None)
        return f"Volume {new:.0f}%."
    if action in ("mute", "unmute"):
        pyautogui.press("volumemute")
        return "Toggled mute."
    if action == "set" and level is not None:
        pyautogui.press("volumedown", presses=50, interval=0)
        pyautogui.press("volumeup", presses=int(round(level / 2)), interval=0)
        return f"Volume {level:.0f}%."
    steps = int(round((level or 10) / 2))
    pyautogui.press("volumeup" if action == "up" else "volumedown", presses=steps, interval=0)
    return f"Volume {action}."


@tool("Set screen brightness (0-100). Works on laptops/supported monitors.", {"level": N("0-100")}, ["level"])
def brightness(level: float):
    code, _, err = run_ps(f"(Get-WmiObject -Namespace root/WMI -Class WmiMonitorBrightnessMethods)"
                          f".WmiSetBrightness(1,{int(level)})", 15)
    if code != 0:
        return "This monitor doesn't support software brightness control; use the monitor's buttons."
    return f"Brightness {int(level)}%."


@tool("Switch Windows between dark mode and light mode.", {"mode": S("Theme", ["dark", "light"])}, ["mode"])
def set_theme(mode: str):
    import winreg
    val = 0 if mode == "dark" else 1
    with winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize",
                        0, winreg.KEY_SET_VALUE) as k:
        winreg.SetValueEx(k, "AppsUseLightTheme", 0, winreg.REG_DWORD, val)
        winreg.SetValueEx(k, "SystemUsesLightTheme", 0, winreg.REG_DWORD, val)
    user32.SendMessageTimeoutW(0xFFFF, 0x001A, 0, "ImmersiveColorSet", 2, 1000, None)
    return f"{mode.capitalize()} mode on."


_SETTINGS = {
    "display": "display", "sound": "sound", "bluetooth": "bluetooth", "wifi": "network-wifi", "network": "network",
    "apps": "appsfeatures", "installed apps": "appsfeatures", "update": "windowsupdate", "windows update": "windowsupdate",
    "battery": "batterysaver", "power": "powersleep", "notifications": "notifications", "personalization": "personalization",
    "background": "personalization-background", "wallpaper": "personalization-background", "themes": "themes",
    "colors": "colors", "mouse": "mousetouchpad", "keyboard": "typing", "privacy": "privacy", "storage": "storagesense",
    "about": "about", "default apps": "defaultapps", "date and time": "dateandtime", "language": "regionlanguage",
    "accounts": "yourinfo", "vpn": "network-vpn", "printers": "printers", "startup apps": "startupapps",
    "security": "windowsdefender", "windows security": "windowsdefender", "night light": "nightlight",
    "focus": "focus", "microphone": "privacy-microphone", "camera": "privacy-webcam", "gaming": "gaming-gamebar",
    "multitasking": "multitasking", "taskbar": "taskbar", "hotspot": "network-mobilehotspot", "proxy": "network-proxy",
}


@tool("Open a Windows Settings page.", {"page": S("Page", list(_SETTINGS))}, ["page"])
def open_settings(page: str):
    os.startfile("ms-settings:" + _SETTINGS.get(page.lower(), page.lower().replace(" ", "")))
    return f"Opened {page} settings."


def _power_q(a):
    act = a.get("action")
    return {"shutdown": "Shut down the PC?", "restart": "Restart the PC?", "sign_out": "Sign out of Windows?",
            "sleep": "Put the PC to sleep?", "hibernate": "Hibernate the PC?"}.get(act)


@tool("Power & session: lock, sleep, hibernate, shutdown, restart, sign_out, screen_off, cancel_shutdown.",
      {"action": S("Action", ["lock", "sleep", "hibernate", "shutdown", "restart", "sign_out", "screen_off",
                              "cancel_shutdown"])}, ["action"], confirm=_power_q)
def system_power(action: str):
    if action == "lock":
        user32.LockWorkStation()
    elif action == "screen_off":
        user32.PostMessageW(0xFFFF, 0x0112, 0xF170, 2)
    elif action == "sleep":
        threading.Timer(1.5, lambda: run_ps("Add-Type -AssemblyName System.Windows.Forms; "
                        "[System.Windows.Forms.Application]::SetSuspendState('Suspend',$false,$false)")).start()
    elif action == "hibernate":
        subprocess.Popen(["shutdown", "/h"], creationflags=NO_WINDOW)
    elif action == "shutdown":
        subprocess.Popen(["shutdown", "/s", "/t", "5"], creationflags=NO_WINDOW)
    elif action == "restart":
        subprocess.Popen(["shutdown", "/r", "/t", "5"], creationflags=NO_WINDOW)
    elif action == "sign_out":
        subprocess.Popen(["shutdown", "/l"], creationflags=NO_WINDOW)
    elif action == "cancel_shutdown":
        subprocess.run(["shutdown", "/a"], capture_output=True, creationflags=NO_WINDOW)
    return {"lock": "Locked.", "shutdown": "Shutting down in 5 seconds.", "restart": "Restarting in 5 seconds.",
            "cancel_shutdown": "Shutdown cancelled."}.get(action, f"{action.replace('_', ' ').capitalize()} done.")


# =============================================================================
# System info, security
# =============================================================================
def _gpu():
    try:
        out = subprocess.run(["nvidia-smi", "--query-gpu=name,utilization.gpu,memory.used,memory.total,temperature.gpu",
                              "--format=csv,noheader,nounits"], capture_output=True, text=True, timeout=3,
                             creationflags=NO_WINDOW).stdout.strip().split(", ")
        return f"GPU {out[0]}: {out[1]}% load, {int(out[2]) / 1024:.1f}/{int(out[3]) / 1024:.0f} GB VRAM, {out[4]}°C"
    except Exception:
        return None


@tool("PC status. detail: overview (CPU/RAM/GPU/disk/battery/uptime), processes (top resource users), "
      "network (IP, Wi-Fi, ping).", {"detail": S("What to report", ["overview", "processes", "network"])})
def system_status(detail: str = "overview"):
    if detail == "processes":
        procs = list(psutil.process_iter(["name", "memory_info"]))
        for p in procs:
            try:
                p.cpu_percent(None)
            except Exception:
                pass
        time.sleep(0.6)
        rows = []
        for p in procs:
            try:
                rows.append((p.cpu_percent(None) / psutil.cpu_count(), p.info["memory_info"].rss / 1e6, p.info["name"]))
            except Exception:
                pass
        agg = {}
        for c, m, n in rows:
            if n in ("System Idle Process", "Idle", "System", "Registry", "Memory Compression"):
                continue
            a = agg.setdefault(n, [0, 0, 0])
            a[0] += c
            a[1] += m
            a[2] += 1
        by_cpu = sorted(agg.items(), key=lambda x: -x[1][0])[:6]
        by_mem = sorted(agg.items(), key=lambda x: -x[1][1])[:6]
        return ("Top CPU: " + ", ".join(f"{n} {v[0]:.0f}%" for n, v in by_cpu) +
                "\nTop memory: " + ", ".join(f"{n} {v[1] / 1024:.1f} GB" for n, v in by_mem))
    if detail == "network":
        parts = []
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            s.connect(("8.8.8.8", 80))
            parts.append(f"Local IP {s.getsockname()[0]}")
            s.close()
        except Exception:
            parts.append("No network connection")
        out = subprocess.run(["netsh", "wlan", "show", "interfaces"], capture_output=True, text=True,
                             creationflags=NO_WINDOW).stdout
        ssid = re.search(r"^\s*SSID\s*:\s*(.+)$", out, re.M)
        sig = re.search(r"^\s*Signal\s*:\s*(.+)$", out, re.M)
        if ssid:
            parts.append(f"Wi-Fi {ssid.group(1).strip()} ({sig.group(1).strip() if sig else '?'})")
        p = subprocess.run(["ping", "-n", "2", "-w", "1500", "1.1.1.1"], capture_output=True, text=True,
                           creationflags=NO_WINDOW).stdout
        avg = re.search(r"Average = (\d+)ms", p)
        parts.append(f"Internet ping {avg.group(1)} ms" if avg else "Internet unreachable")
        return ", ".join(parts)
    vm = psutil.virtual_memory()
    parts = [dt.datetime.now().strftime("%A %d %B %Y, %I:%M %p"),
             f"CPU {psutil.cpu_percent(0.4):.0f}%", f"RAM {vm.used / 1e9:.1f}/{vm.total / 1e9:.0f} GB ({vm.percent:.0f}%)"]
    g = _gpu()
    if g:
        parts.append(g)
    for part in psutil.disk_partitions():
        try:
            u = psutil.disk_usage(part.mountpoint)
            parts.append(f"{part.device[:2]} {u.free / 1e9:.0f} GB free of {u.total / 1e9:.0f}")
        except Exception:
            pass
    b = psutil.sensors_battery()
    if b:
        parts.append(f"Battery {b.percent:.0f}% {'charging' if b.power_plugged else 'on battery'}")
    up = time.time() - psutil.boot_time()
    parts.append(f"Up {int(up // 86400)}d {int(up % 86400 // 3600)}h")
    return ". ".join(parts)


def _kill_q(a):
    return f"End the process {a.get('name')}?"


@tool("Force-end a process by name (e.g. a frozen app).", {"name": S("Process name")}, ["name"], confirm=_kill_q)
def kill_process(name: str):
    return close_app(name, force=True)


@tool("Security: 'scan' checks the PC for threats & security problems (Defender, firewall, suspicious "
      "programs, open ports, updates) and opens a report; 'quick_virus_scan'/'full_virus_scan' run Defender; "
      "'update_definitions' updates antivirus.",
      {"action": S("Action", ["scan", "quick_virus_scan", "full_virus_scan", "update_definitions"])})
def security_check(action: str = "scan"):
    if action == "scan":
        CTX.status("Scanning for threats…")
        r = security.scan()
        security.open_report(r["report"])
        return security.summary(r)
    if action == "update_definitions":
        return security.update_definitions()
    return security.virus_scan("full" if action.startswith("full") else "quick", CTX.announce)


def _clean_q(a):
    return {"recycle_bin": "Permanently empty the recycle bin?",
            "temp_files": "Delete temporary files to free up space?"}.get(a.get("target"))


@tool("Free disk space: delete temp files or empty the recycle bin.",
      {"target": S("What to clean", ["temp_files", "recycle_bin"])}, ["target"], confirm=_clean_q)
def clean_up(target: str):
    if target == "recycle_bin":
        ctypes.windll.shell32.SHEmptyRecycleBinW(None, None, 0x7)
        return "Recycle bin emptied."
    freed, cutoff = 0, time.time() - 86400
    for root in {os.environ.get("TEMP", ""), r"C:\Windows\Temp"}:
        for dirpath, dirs, files in os.walk(root):
            for f in files:
                p = os.path.join(dirpath, f)
                try:
                    st = os.stat(p)
                    if st.st_mtime < cutoff:
                        os.remove(p)
                        freed += st.st_size
                except Exception:
                    pass
    return f"Freed {freed / 1e6:.0f} MB of temporary files."


# =============================================================================
# Keyboard, mouse, screen
# =============================================================================
@tool("Type text into the focused app (like the user typing).",
      {"text": S("Text to type"), "press_enter": B("Press Enter after")}, ["text"])
def type_text(text: str, press_enter: bool = False):
    import pyautogui
    import pyperclip
    try:
        old = pyperclip.paste()
    except Exception:
        old = None
    pyperclip.copy(text)
    time.sleep(0.05)
    pyautogui.hotkey("ctrl", "v")
    time.sleep(0.15)
    if press_enter:
        pyautogui.press("enter")
    if old is not None:
        threading.Timer(0.6, lambda: pyperclip.copy(old)).start()
    return f"Typed {len(text)} characters."


_KEYMAP = {"windows": "win", "window": "win", "start": "win", "control": "ctrl", "escape": "esc", "return": "enter",
           "del": "delete", "page up": "pageup", "page down": "pagedown", "arrow up": "up", "arrow down": "down",
           "arrow left": "left", "arrow right": "right", "space bar": "space", "spacebar": "space",
           "caps lock": "capslock", "print screen": "printscreen", "cmd": "win", "super": "win", "option": "alt"}


@tool("Press keys or shortcuts, e.g. 'ctrl+c', 'alt+tab', 'win+d', 'enter', 'ctrl+shift+esc'. "
      "Comma-separate for a sequence.", {"keys": S("Key combo(s)"), "times": N("Repeat count")}, ["keys"])
def press_keys(keys: str, times: float = 1):
    import pyautogui
    for _ in range(max(1, min(int(times), 50))):
        for combo in [c for c in re.split(r"\s*,\s*", keys.strip().lower()) if c]:
            parts = [_KEYMAP.get(k.strip(), k.strip()) for k in re.split(r"\s*\+\s*", combo)]
            bad = [k for k in parts if k not in pyautogui.KEYBOARD_KEYS]
            if bad:
                return f"Unknown key(s): {', '.join(bad)}"
            pyautogui.hotkey(*parts) if len(parts) > 1 else pyautogui.press(parts[0])
            time.sleep(0.05)
    return f"Pressed {keys}."


@tool("Mouse: move, click, double_click, right_click, scroll_up, scroll_down, position. Coordinates are screen pixels.",
      {"action": S("Mouse action", ["move", "click", "double_click", "right_click", "scroll_up", "scroll_down",
                                    "position"]),
       "x": N("X pixel"), "y": N("Y pixel"), "amount": N("Scroll amount (default 5)")}, ["action"])
def mouse(action: str, x: float = None, y: float = None, amount: float = 5):
    import pyautogui
    pyautogui.FAILSAFE = False
    pos = (int(x), int(y)) if x is not None and y is not None else None
    if action == "position":
        p = pyautogui.position()
        return f"Mouse at {p.x}, {p.y}. Screen is {pyautogui.size().width}x{pyautogui.size().height}."
    if action == "move" and pos:
        pyautogui.moveTo(*pos, duration=0.15)
    elif action == "click":
        pyautogui.click(*(pos or ()))
    elif action == "double_click":
        pyautogui.doubleClick(*(pos or ()))
    elif action == "right_click":
        pyautogui.rightClick(*(pos or ()))
    elif action in ("scroll_up", "scroll_down"):
        pyautogui.scroll(int(amount * 120) * (1 if action == "scroll_up" else -1), *(pos or ()))
    return f"Mouse {action.replace('_', ' ')} done."


@tool("Read the buttons, links, fields and text in the active window, to know what can be clicked.",
      {"include_text": B("Also include plain text on screen")})
def read_screen(include_text: bool = True):
    return uia.list_elements(include_text)


@tool("Click/press/type into an element of the active window by its visible name, e.g. 'Send', 'Sign in', "
      "'Search'. Use read_screen first if unsure of the name.",
      {"name": S("Element's visible name"), "action": S("Action", ["click", "double_click", "right_click", "type",
                                                                    "focus"]),
       "text": S("Text to type when action=type"), "kind": S("Optional type, e.g. button, edit, link")}, ["name"])
def click_element(name: str, action: str = "click", text: str = "", kind: str = ""):
    return uia.click_element(name, kind, action, text)


@tool("Take a screenshot and save it to Pictures\\Screenshots.")
def take_screenshot():
    import pyautogui
    folder = (known_folder("pictures") or Path.home()) / "Screenshots"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"Miles_{dt.datetime.now():%Y%m%d_%H%M%S}.png"
    pyautogui.screenshot().save(path)
    return f"Screenshot saved to {path}"


def _screenshot_b64(max_side=1280):
    import pyautogui
    img = pyautogui.screenshot()
    full = img.size
    img.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=85)
    return base64.b64encode(buf.getvalue()).decode(), full, img.size


def vision_chat(prompt: str, b64: str) -> str:
    """Ask the vision model about an image. When the main brain can see, it answers itself (no model swap)."""
    vm = CTX.vision_model
    same = vm == config.OLLAMA_MODEL
    if same and CTX.side_call:          # inside the running conversation: reuses the cached prompt (much faster)
        return CTX.side_call(prompt, b64)
    body = {"model": vm, "stream": False, "messages": [{"role": "user", "content": prompt, "images": [b64]}],
            # a separate vision model can't share 12 GB of VRAM with the main one: unload it right after
            "keep_alive": config.OLLAMA_KEEP_ALIVE if same else 0}
    if same:
        body["options"] = {"num_ctx": config.OLLAMA_CONTEXT}   # same options = no model reload
    for think in (False, None):
        if think is not None:
            body["think"] = think
        else:
            body.pop("think", None)
        r = requests.post(f"{config.OLLAMA_URL}/api/chat", timeout=180, json=body)
        if r.status_code == 400 and "think" in r.text:
            continue
        r.raise_for_status()
        return re.sub(r"<think>.*?</think>", "", r.json()["message"]["content"], flags=re.S).strip()
    return ""


@tool("Read all text visible in the active window (Windows OCR) to answer questions about the screen or check "
      "a result.", {"question": S("What to find out")}, ["question"])
def look_at_screen(question: str = ""):
    """Lite: no vision model - Windows OCR reads the window's text in ~0.3 s on a laptop CPU."""
    from . import ocr
    CTX.status("Reading the screen…")
    _, title = _foreground_title()
    try:
        text = ocr.screen_text(active_window_only=True, max_chars=2500)
    except Exception as e:
        return (f"FAILED: the screen can't be captured right now ({e}); it may be locked or asleep. Tell the "
                f"user that plainly.")
    return (f"Active window: '{title}'. Its visible text, top to bottom (from OCR, may have small errors):\n"
            f"{text or '(no readable text - it may be an image or video)'}\n---\nQuestion: {question}\n"
            f"Answer in one or two spoken sentences (what the window is and what matters in it). Do NOT read "
            f"the text out.")


@tool("Click visible text anywhere on screen, e.g. 'Sign in', 'Next', 'Accept all' (OCR). Use when "
      "click_element can't find it.", {"target": S("The text to click"),
                                       "action": S("Mouse action", ["click", "double_click", "right_click", "hover"])},
      ["target"])
def click_on_screen(target: str, action: str = "click"):
    import pyautogui
    from . import ocr
    CTX.status(f"Looking for '{target}'…")
    pos, found, score = ocr.find(target)
    if not pos or score < 0.75:
        return f"I can't see the text '{target}' on screen." + (f" Closest: '{found}'." if found else "")
    pyautogui.FAILSAFE = False
    {"double_click": pyautogui.doubleClick, "right_click": pyautogui.rightClick,
     "hover": pyautogui.moveTo}.get(action, pyautogui.click)(*pos)
    return f"{action.replace('_', ' ').capitalize()}ed '{found}' at {pos}."


@tool("Wait a few seconds (e.g. for an app/page to load between steps).", {"seconds": N("Seconds")}, ["seconds"])
def wait(seconds: float):
    time.sleep(max(0.1, min(float(seconds), 15)))
    return f"Waited {seconds} s."


# =============================================================================
# Files
# =============================================================================
@tool("Find files or folders by name on the PC (Desktop, Downloads, Documents, etc. or a given folder/drive).",
      {"name": S("Part of the file/folder name"), "folder": S("Optional folder or drive to search, e.g. 'D:'")},
      ["name"])
def find_files(name: str, folder: str = None):
    CTX.status(f"Searching for {name}…")
    roots = [resolve_path(folder)] if folder else user_search_roots()
    q = name.lower().strip().strip("*")
    words = q.split()
    results, seen, t0 = [], set(), time.time()
    skip = {"appdata", "node_modules", ".git", "$recycle.bin", "windows", "program files", "program files (x86)",
            "__pycache__", ".venv", "venv", "site-packages", ".cache"}
    for root in roots:
        for dirpath, dirs, files in os.walk(root):
            dirs[:] = [d for d in dirs if d.lower() not in skip and not d.startswith(".")]
            for n in dirs + files:
                nl = n.lower()
                if all(w in nl for w in words):
                    p = os.path.join(dirpath, n)
                    if p.lower() not in seen:
                        seen.add(p.lower())
                        results.append(p)
            if len(results) >= 20 or time.time() - t0 > 8:
                break
        if len(results) >= 20 or time.time() - t0 > 8:
            break
    if not results:
        return f"Nothing named '{name}' found." + ("" if folder else " Try giving a folder or drive.")
    results.sort(key=lambda p: (not os.path.basename(p).lower().startswith(q), len(p)))
    return "Found:\n" + "\n".join(results[:15])


def _paths(path: str):
    return [p.strip() for p in str(path).split("|") if p.strip()]


def _file_q(a):
    act, paths = a.get("action"), _paths(a.get("path", ""))
    what = Path(paths[0]).name if len(paths) == 1 else f"these {len(paths)} items"
    if act == "delete":
        return f"Move {what} to the recycle bin?"
    if act == "move":
        return f"Move {what} to {a.get('destination')}?"
    if act == "write" and paths and resolve_path(paths[0]).exists():
        return f"Overwrite {what}?"
    return None


@tool("File operations: open (file/folder/'downloads' etc.), list, read, write, append, create_folder, copy, "
      "move (also rename), delete (to recycle bin), zip. For several files (e.g. selected in Explorer) join the "
      "paths with ' | '.",
      {"action": S("Operation", ["open", "list", "read", "write", "append", "create_folder", "copy", "move",
                                 "delete", "zip"]),
       "path": S("File/folder path(s); may start with desktop/, downloads/, documents/ etc."),
       "destination": S("Target for copy/move, or the .zip to create"), "content": S("Text for write/append")},
      ["action", "path"], confirm=_file_q)
def file_action(action: str, path: str, destination: str = None, content: str = ""):
    paths = _paths(path)
    if len(paths) > 1 or action == "zip":
        return _many_files(action, [resolve_path(p) for p in paths], destination)
    p = resolve_path(path)
    if action == "open":
        if not p.exists():
            return f"{p} doesn't exist. Use find_files to locate it."
        os.startfile(p)
        return f"Opened {p}"
    if action == "list":
        if not p.is_dir():
            return f"{p} is not a folder."
        items = sorted(p.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower()))
        return f"{p} ({len(items)} items):\n" + "\n".join(
            f"{'[folder] ' if i.is_dir() else ''}{i.name}" for i in items[:60])
    if action == "read":
        text = p.read_text(encoding="utf-8", errors="replace")
        return text[:4000] + ("\n…(truncated)" if len(text) > 4000 else "")
    if action in ("write", "append"):
        p.parent.mkdir(parents=True, exist_ok=True)
        with open(p, "a" if action == "append" else "w", encoding="utf-8") as f:
            f.write(content)
        return f"Saved {p}"
    if action == "create_folder":
        p.mkdir(parents=True, exist_ok=True)
        return f"Created {p}"
    if action in ("copy", "move"):
        if not destination:
            return "Need a destination."
        d = resolve_path(destination)
        if d.is_dir():
            d = d / p.name
        elif "\\" not in destination and "/" not in destination and not re.match(r"[a-zA-Z]:", destination):
            d = p.parent / destination                     # plain new name = rename in place
        if action == "copy":
            (shutil.copytree if p.is_dir() else shutil.copy2)(p, d)
        else:
            shutil.move(str(p), str(d))
        return f"{action.capitalize()}d to {d}"
    if action == "delete":
        from send2trash import send2trash
        send2trash(str(p))
        return f"Moved {p.name} to the recycle bin."
    return "Unknown action."


def _many_files(action: str, paths: list, destination: str = None) -> str:
    missing = [p.name for p in paths if not p.exists()]
    if missing:
        return f"Not found: {', '.join(missing[:5])}. Nothing was changed."
    if action == "zip":
        import zipfile
        dest = resolve_path(destination) if destination else \
            paths[0].parent / f"{paths[0].stem if len(paths) == 1 else paths[0].parent.name or 'Files'}.zip"
        if dest.suffix.lower() != ".zip":
            dest = dest.with_suffix(".zip")
        CTX.status(f"Zipping {len(paths)} item(s)…")
        with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
            for p in paths:
                if p.is_dir():
                    for f in p.rglob("*"):
                        z.write(f, f.relative_to(p.parent))
                else:
                    z.write(p, p.name)
        return f"Zipped {len(paths)} item(s) into {dest}"
    if action in ("copy", "move"):
        if not destination:
            return "Need a destination folder."
        d = resolve_path(destination)
        d.mkdir(parents=True, exist_ok=True)
        for p in paths:
            if action == "copy":
                (shutil.copytree if p.is_dir() else shutil.copy2)(p, d / p.name)
            else:
                shutil.move(str(p), str(d / p.name))
        return f"{action.capitalize()}d {len(paths)} items to {d}"
    if action == "delete":
        from send2trash import send2trash
        for p in paths:
            send2trash(str(p))
        return f"Moved {len(paths)} items to the recycle bin."
    if action == "open":
        for p in paths[:10]:
            os.startfile(p)
        return f"Opened {min(len(paths), 10)} items."
    return f"'{action}' works on one file at a time."


# =============================================================================
# Productivity
# =============================================================================
@tool("Get or set the clipboard text.", {"action": S("get or set", ["get", "set"]), "text": S("Text to copy")},
      ["action"])
def clipboard(action: str, text: str = ""):
    import pyperclip
    if action == "set":
        pyperclip.copy(text)
        return "Copied to clipboard."
    return "Clipboard: " + (pyperclip.paste() or "(empty)")[:2000]


_REMINDERS = []


def _parse_clock(s: str):
    s = s.strip().lower().replace(".", "")
    for fmt in ("%H:%M", "%I:%M %p", "%I:%M%p", "%I %p", "%I%p"):
        try:
            t = dt.datetime.strptime(s, fmt).time()
            when = dt.datetime.combine(dt.date.today(), t)
            return when + dt.timedelta(days=1) if when < dt.datetime.now() else when
        except ValueError:
            pass
    return None


@tool("Set a timer/reminder, or SCHEDULE A COMMAND to run later (e.g. at 18:00 'play lofi on Spotify', in 30 "
      "minutes 'shut down the PC'). Give minutes/seconds from now, or at_time like '17:30' / '5:30 pm'.",
      {"message": S("What to say when it fires"), "minutes": N("Minutes from now"), "seconds": N("Seconds from now"),
       "at_time": S("Clock time"), "command": S("Optional plain-English command to execute at that time")},
      ["message"])
def set_reminder(message: str, minutes: float = 0, seconds: float = 0, at_time: str = None, command: str = ""):
    if at_time:
        when = _parse_clock(at_time)
        if not when:
            return f"Couldn't understand the time '{at_time}'."
        delay = (when - dt.datetime.now()).total_seconds()
    else:
        delay = float(minutes) * 60 + float(seconds)
        when = dt.datetime.now() + dt.timedelta(seconds=delay)
    if delay <= 0:
        return "Give me a time in the future."

    def fire():
        if command:
            CTX.announce(f"{message}." if message else f"Running your scheduled task, {config.USER_TITLE}.")
            CTX.run_steps([command])
        else:
            CTX.announce(f"Reminder, {config.USER_TITLE}: {message}.")
        _REMINDERS[:] = [r for r in _REMINDERS if r[0] != when]
    t = threading.Timer(delay, fire)
    t.daemon = True
    t.start()
    _REMINDERS.append((when, message or command))
    what = f"I'll {command}" if command else message
    return f"Scheduled for {when:%I:%M %p}: {what}."


@tool("Current weather and today's forecast.", {"location": S("City (blank = current location)")})
def weather(location: str = ""):
    loc = urllib.parse.quote(location or config.HOME_CITY or "")
    d = requests.get(f"https://wttr.in/{loc}?format=j1", timeout=8).json()
    c = d["current_condition"][0]
    today = d["weather"][0]
    area = d.get("nearest_area", [{}])[0].get("areaName", [{"value": location or "here"}])[0]["value"]
    rain = max(int(h.get("chanceofrain", 0)) for h in today["hourly"])
    return (f"{area}: {c['weatherDesc'][0]['value']}, {c['temp_C']}°C (feels {c['FeelsLikeC']}°C), humidity "
            f"{c['humidity']}%. Today {today['mintempC']}–{today['maxtempC']}°C, {rain}% chance of rain.")


# =============================================================================
# Memory, asking the user, writing, planning
# =============================================================================
def _forget_q(a):
    if a.get("forget") and _norm(a.get("fact", "")) in ("everything", "all", "all facts", "everything about me"):
        return "Erase everything I remember about you?"
    return None


@tool("Long-term memory: save a lasting fact about the user (name, likes/dislikes, people and their details, "
      "routines, important dates) as one short sentence, e.g. \"Arnav's sister is Priya\". Saving a newer version "
      "of a fact updates it. forget=true removes matching facts.",
      {"fact": S("The fact, as a short self-contained sentence"), "forget": B("Remove matching facts instead")},
      ["fact"], confirm=_forget_q)
def remember(fact: str, forget: bool = False):
    from . import memory
    if forget:
        gone = memory.forget(fact)
        return f"Forgot: {'; '.join(gone)}." if gone else "Nothing in memory matched that."
    return memory.add(fact)


@tool("Search your memory of past conversations and saved facts: 'what did I ask you yesterday', 'the song you "
      "played earlier', 'when did I email Tulsi', 'what do you know about my exams'.",
      {"query": S("Keywords to look for (can be empty when 'when' is given)"),
       "when": S("Optional time filter: today, yesterday, this week, this month, or a date")}, ["query"])
def recall(query: str = "", when: str = ""):
    from . import memory
    return memory.recall(query, when)


@tool("Ask the user to TYPE a detail you don't know or didn't hear clearly (an email address, a spelling, which "
      "file). Opens a box on screen; better than guessing. Never for passwords or card numbers.",
      {"question": S("Short question"), "suggestion": S("Your best guess, pre-filled"),
       "remember_as": S("Label to save the answer in memory, e.g. \"Tulsi's phone number\"")}, ["question"])
def ask_user(question: str, suggestion: str = "", remember_as: str = ""):
    from .mail import _PLACEHOLDER
    if _PLACEHOLDER.search(suggestion or "") or re.search(r"\b(example|placeholder|xxx)\b", suggestion or "", re.I):
        suggestion = ""                              # never pre-fill a made-up guess
    answer = CTX.ask(question, suggestion)
    if answer is None:
        return ("The user skipped the box. Don't ask again. Whatever needed this answer was NOT done - say so "
                "plainly; don't claim it's ready or done.")
    answer = answer.strip()
    if not answer:
        return "The user left it empty."
    note = ""
    if remember_as:
        from . import mail, memory
        who = re.match(r"^(?:my |the )?(.+?)'s? e-?mail", remember_as.strip(), re.I)
        addr = mail.spoken_to_email(answer)
        if who and mail.EMAIL_RE.match(addr):          # an email address: it belongs in contacts
            mail.save_contact(who.group(1).strip().title() if who.group(1).islower() else who.group(1).strip(),
                              addr)
            note = f" (saved to contacts as {who.group(1).strip()})"
        else:
            memory.add(f"{remember_as.rstrip(': ')}: {answer}")
            note = " (saved to memory)"
    return f"The user typed: {answer}{note}"


@tool("WRITE text for the user (note, message, reply, poem, essay, list...) or a refined/rewritten/translated "
      "version of selected text. Then write the text itself as your next message: it's typed out, not spoken. "
      "where: auto = the text box they're in, else Notepad; selection = replace the highlighted text; replace = "
      "replace the whole box; notepad; word; clipboard. Emails: compose_email. Exact words: type_text.",
      {"request": S("What to write, close to the user's words"),
       "where": S("Where it goes", ["auto", "selection", "replace", "notepad", "word", "clipboard"])},
      ["request"])
def write(request: str, where: str = "auto"):
    from .writer import Writer
    w = Writer(where, request)
    place = w.prepare()
    CTX.writer = w
    CTX.status(f"Writing in {place}…")
    # Lite: short and literal (a laptop CPU reads ~50 tokens/s, and a small model copies examples/instructions)
    snap = CTX.snapshot
    original = ""
    if w.where == "selection" and snap is not None and snap.selection:
        original = snap.selection[:2000]
    elif w.where == "replace" and snap is not None and snap.field_text:
        original = snap.field_text[:2000]
    # The brain sends this as a user message right before the writing turn: a small model follows a direct
    # instruction far better than one buried in a tool result.
    if original:
        w.prompt = (f'{request}\n\nText:\n"""{original}"""\n\nRewrite the text as asked - clearly better, not the '
                    f"same words. Reply with only the new text.")
    else:
        w.prompt = f"{request}\n\nReply with only the finished text - no introduction, no quotes."
    return f"Ready: your next message is typed into {place}."


@tool("Read the user's current context: selected text, the focused text box and its text, page address, files "
      "selected in Explorer, what the mouse is over, optionally the clipboard. Use if [Context] lacks it.",
      {"clipboard": B("Also read the clipboard")})
def get_context(clipboard: bool = False):
    from . import context
    s = context.snapshot(selection=True, clipboard=clipboard)
    CTX.snapshot = s
    return f"Window: '{s.title}' ({s.proc}). " + (s.details() or "Nothing is selected and no text box is focused.")


@tool("Change how you speak (remembered): faster, slower, a speed (0.8-1.6), or a voice: george, lewis, daniel, "
      "fable (British men), michael, adam (American men), emma, isabella (British women), heart, bella (American).",
      {"action": S("What to change", ["faster", "slower", "speed", "voice", "reset"]),
       "speed": N("Speed"), "voice": S("Voice name")}, ["action"])
def voice_settings(action: str, speed: float = None, voice: str = ""):
    from . import persona
    return persona.voice_setting(action, speed, voice)


@tool("For long or multi-part requests: list every step in order BEFORE starting, then carry out all of them.",
      {"steps": {"type": "array", "items": {"type": "string"}, "description": "The steps, in order"}}, ["steps"])
def plan(steps: list):
    steps = [str(s).strip() for s in (steps or []) if str(s).strip()][:20]
    if not steps:
        return "No steps given."
    CTX.plan = steps
    CTX.status("Plan: " + "  →  ".join(steps))
    return f"Plan noted ({len(steps)} steps). Now do step 1: {steps[0]}"


_READ_ONLY = re.compile(r"^\s*\(?\s*(get-|test-|select-|measure-|format-|where-|sort-|group-|write-output|"
                        r"resolve-|convertto-json|\$env:|\[environment\]|systeminfo|ipconfig|hostname|whoami|ping|"
                        r"tasklist|nslookup|echo|dir\b|ls\b|cat\b|type\b|winget (search|list|show))", re.I)
_MUTATING = re.compile(r"(remove-|set-|stop-|start-|new-|rename-|move-|copy-|clear-|invoke-|restart-|out-file|"
                       r"add-|disable-|enable-|install-|uninstall-|update-|format-volume|\bdel\b|\brm\b|\brd\b|"
                       r"reg\s+(add|delete)|shutdown|taskkill|>\s*\S)", re.I)


def _cmd_q(a):
    cmd = a.get("command", "")
    if config.AUTO_APPROVE_READ_ONLY and _READ_ONLY.match(cmd) and not _MUTATING.search(cmd):
        return None                                   # read-only: no need to bother the user
    return f"I'm about to {a.get('explanation') or 'run a system command'}. Shall I proceed?"


@tool("Run a PowerShell command for anything the other tools can't do (system info, settings, files, "
      "processes, network, registry...). Returns its output. Read-only commands run immediately.",
      {"command": S("PowerShell command"), "explanation": S("Plain-English description of what it does")},
      ["command", "explanation"], confirm=_cmd_q)
def run_command(command: str, explanation: str = ""):
    CTX.status(f"> {command[:80]}")
    try:
        code, out, err = run_ps(command, config.COMMAND_TIMEOUT_SEC)
    except subprocess.TimeoutExpired:
        return "Command timed out."
    res = (out + ("\n" + err if err else "")).strip() or f"(no output, exit code {code})"
    return res[:3000]


# =============================================================================
# Browser tabs, Spotify, email, contacts
# =============================================================================
@tool("Browser TABS (Brave/Chrome/Edge): close/switch/mute/unmute/reload ONE tab by its title (or the current "
      "tab if no title), open a new tab, reopen the last closed tab, next/previous, list. ALWAYS use this for "
      "tabs - never close_app or manage_window, those close the whole browser.",
      {"action": S("Tab action", ["close", "switch", "mute", "unmute", "reload", "new", "reopen", "next",
                                  "previous", "list"]),
       "title": S("Part of the tab's title, e.g. 'YouTube'. Omit for the current tab.")}, ["action"])
def browser_tab(action: str, title: str = ""):
    from . import browser
    return browser.tab_action(action, title)


@tool("Spotify: play a song/artist/album/playlist (searches and plays the top result), resume, pause, next, "
      "previous, shuffle, repeat, now_playing. Use this for ALL music unless the user says YouTube.",
      {"action": S("Action", ["play", "resume", "pause", "next", "previous", "shuffle", "repeat", "now_playing"]),
       "query": S("What to play, for action=play")}, ["action"])
def spotify(action: str, query: str = ""):
    from . import spotify as sp
    if action == "play":
        if not query or _norm(query) in ("something", "anything", "music", "some music", "my music", "a song"):
            return sp.control("resume")
        return sp.play(query)
    if action == "now_playing":
        return sp.now_playing()
    return sp.control(action)


@tool("Write an email and open it in Gmail ready to go. YOU write a complete subject and a polished body (greeting, "
      "message, sign-off) from the user's gist. send=true only if the user explicitly said to send it; they "
      "will be asked to confirm. Unknown or unclear addresses are asked for in a box and saved automatically.",
      {"to": S("Recipient names or addresses, comma-separated; 'Name <address>' when you have both"), "subject": S("Subject line"),
       "body": S("Full email body"), "cc": S("Optional CC"), "send": B("Send right away (after confirmation)")},
      ["to", "subject", "body"])
def compose_email(to: str, subject: str, body: str, cc: str = "", send: bool = False):
    from . import mail
    signer = config.USER_NAME.strip().title() if config.USER_NAME.strip().islower() else config.USER_NAME.strip()
    body = re.sub(r"\n\s*(sir|boss|\[your name\]|your name)\s*$", f"\n{signer}" if signer else "", body.rstrip(),
                  flags=re.I)                        # it calls the user "sir" - but mustn't sign their emails so
    if signer and signer.lower() not in body[-80:].lower():
        body = body.rstrip() + f"\n\n{signer}"
    return mail.compose(to, subject, body, cc, send, confirm=CTX.confirm, status=CTX.status, ask=CTX.ask,
                        heard=CTX.input_mode == "voice", request=CTX.request_text,
                        trusted={mail.spoken_to_email(a) for a in CTX.typed_answers})


@tool("Send the email currently open in a compose window (the user is asked to confirm first).")
def send_email():
    from . import mail
    if not CTX.confirm("Send the open email now?"):
        return "The user said not to send it."
    return mail.press_send()


@tool("Save or update a contact's email address and/or phone number (so emails/messages can use their name).",
      {"name": S("Contact name"), "email": S("Email address (spoken form is fine)"), "phone": S("Phone number")},
      ["name"])
def save_contact(name: str, email: str = "", phone: str = ""):
    from . import mail
    return mail.save_contact(name, email, phone)


# =============================================================================
# Knowledge: research, reading pages, news, briefing
# =============================================================================
@tool("Look up CURRENT information on the internet (news, prices, scores, releases, anything recent or that "
      "you're not sure about). Returns top results with snippets and URLs.", {"query": S("Search query")},
      ["query"])
def research(query: str):
    from . import web
    CTX.status(f"Researching: {query}")
    res = web.search(query)
    if not res:
        return "No results (the internet may be down)."
    return "\n".join(f"{i}. {t} - {s} ({u})" for i, (t, u, s) in enumerate(res, 1))


@tool("Read the main text of a web page (e.g. a URL from research).", {"url": S("Page URL")}, ["url"])
def read_webpage(url: str):
    from . import web
    CTX.status("Reading the page…")
    return web.read_page(url)


@tool("Latest news headlines, optionally about a topic.", {"topic": S("Optional topic")})
def news(topic: str = ""):
    from . import web
    items = web.news(topic, 6)
    return "Headlines: " + " | ".join(f"{t} ({s})" for t, s in items) if items else "Couldn't fetch the news."


@tool("A spoken daily briefing: time, weather, system health, upcoming reminders and top headlines.")
def briefing():
    now = dt.datetime.now()
    part = "morning" if now.hour < 12 else "afternoon" if now.hour < 17 else "evening"
    lines = [f"Good {part}, {config.USER_TITLE}. It's {now:%I:%M %p}".replace(" 0", " ") + f" on {now:%A}."]
    try:
        from . import web  # noqa
        w = weather("")
        m = re.match(r"(.+?): (.+?), (-?\d+)°C.*?(\d+)% chance of rain", w)
        if m:
            lines.append(f"It's {m.group(3)} degrees and {m.group(2).lower().strip()} in {m.group(1)}, with a "
                         f"{m.group(4)} percent chance of rain.")
    except Exception:
        pass
    issues = []
    vm = psutil.virtual_memory()
    if vm.percent > 85:
        issues.append(f"memory is at {vm.percent:.0f} percent")
    du = psutil.disk_usage(os.environ.get("SystemDrive", "C:") + "\\")
    if du.free < 15e9:
        issues.append(f"the system drive has only {du.free / 1e9:.0f} gigabytes free")
    lines.append("Systems: " + "; ".join(issues) + "." if issues else "All systems are nominal.")
    soon = [(w, m) for w, m in _REMINDERS if w - now < dt.timedelta(hours=24)]
    if soon:
        lines.append("Coming up: " + "; ".join(f"{m} at {w:%I:%M %p}".replace(" 0", " ") for w, m in soon[:3]) + ".")
    try:
        from . import web
        heads = web.news("", 3)
        if heads:
            lines.append("In the news: " + ". ".join(t for t, _ in heads) + ".")
    except Exception:
        pass
    return " ".join(lines)


# =============================================================================
# Documents, folders, installing apps
# =============================================================================
@tool("Create a document (Word .docx by default, or txt/md) with content YOU write, save it in "
      "Documents\\Miles and open it. Use '# ' for headings and '- ' for bullets in content.",
      {"title": S("Document title / file name"), "content": S("Full document text"),
       "format": S("File type", ["docx", "txt", "md"])}, ["title", "content"])
def create_document(title: str, content: str, format: str = "docx"):
    folder = (known_folder("documents") or Path.home()) / "Miles"
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r'[<>:"/\\|?*]+', "", title).strip()[:80] or "Untitled"
    path = folder / f"{safe}.{format}"
    if format == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        for raw in content.split("\n"):
            s = raw.strip()
            if not s:
                continue
            h = re.match(r"^(#{1,3})\s+(.*)", s)
            if h:
                doc.add_heading(h.group(2), len(h.group(1)))
            elif re.match(r"^[-*•]\s+", s):
                doc.add_paragraph(re.sub(r"^[-*•]\s+", "", s), style="List Bullet")
            elif re.match(r"^\d+[.)]\s+", s):
                doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", s), style="List Number")
            else:
                doc.add_paragraph(s)
        doc.save(path)
    else:
        path.write_text((f"# {title}\n\n" if format == "md" else "") + content, encoding="utf-8")
    os.startfile(path)
    return f"Created and opened {path}"


_CATEGORIES = {
    "Images": {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".heic", ".svg", ".ico", ".tiff"},
    "Videos": {".mp4", ".mkv", ".mov", ".avi", ".webm", ".wmv", ".flv"},
    "Music": {".mp3", ".wav", ".flac", ".m4a", ".aac", ".ogg"},
    "Documents": {".pdf", ".doc", ".docx", ".txt", ".rtf", ".odt", ".xls", ".xlsx", ".csv", ".ppt", ".pptx", ".md"},
    "Archives": {".zip", ".rar", ".7z", ".tar", ".gz", ".iso"},
    "Installers": {".exe", ".msi", ".msix", ".appx", ".apk"},
    "Code": {".py", ".js", ".ts", ".html", ".css", ".json", ".java", ".c", ".cpp", ".cs", ".go", ".rs", ".ipynb"},
}


def _organize_q(a):
    p = resolve_path(a.get("path", "downloads"))
    n = sum(1 for f in p.iterdir() if f.is_file()) if p.is_dir() else 0
    return f"Sort the {n} files in {p.name} into folders by type?" if n else None


@tool("Tidy a folder: move its loose files into subfolders by type (Images, Videos, Documents, Installers...).",
      {"path": S("Folder, e.g. 'downloads' or 'desktop'")}, ["path"], confirm=_organize_q)
def organize_folder(path: str):
    p = resolve_path(path)
    if not p.is_dir():
        return f"{p} isn't a folder."
    moved = {}
    for f in list(p.iterdir()):
        if not f.is_file() or f.name.startswith(".") or f.suffix.lower() in (".lnk", ".ini", ".url"):
            continue
        cat = next((c for c, exts in _CATEGORIES.items() if f.suffix.lower() in exts), "Other")
        dest = p / cat
        dest.mkdir(exist_ok=True)
        target = dest / f.name
        i = 1
        while target.exists():
            target = dest / f"{f.stem} ({i}){f.suffix}"
            i += 1
        try:
            shutil.move(str(f), str(target))
            moved[cat] = moved.get(cat, 0) + 1
        except Exception:
            pass
    if not moved:
        return f"{p.name} was already tidy."
    return f"Organized {sum(moved.values())} files in {p.name}: " + ", ".join(f"{n} {c}" for c, n in moved.items())


def _winget_search(name):
    code, out, _ = run_ps(f'winget search --name "{name}" --accept-source-agreements --count 5', 60)
    lines = out.splitlines()
    try:
        start = next(i for i, l in enumerate(lines) if set(l.strip()) == {"-"}) + 1
    except StopIteration:
        return []
    rows = []
    for l in lines[start:]:
        cols = re.split(r"\s{2,}", l.strip())
        if len(cols) >= 2:
            rows.append((cols[0], cols[1]))
    return rows


@tool("Install an application with winget (Windows Package Manager), e.g. 'VLC', 'Zoom', 'Python'. Asks first; "
      "installs in the background and announces when done.", {"name": S("App name")}, ["name"])
def install_app(name: str):
    CTX.status(f"Finding {name}…")
    rows = _winget_search(name)
    if not rows:
        return f"winget couldn't find an app called {name}."
    pkg, pid = rows[0]
    if not CTX.confirm(f"Install {pkg} from the Windows package manager?"):
        return "Installation cancelled."

    def run():
        code, out, err = run_ps(f'winget install --id "{pid}" -e --silent --accept-package-agreements '
                                f'--accept-source-agreements', 1800)
        ok = code == 0 or "successfully installed" in out.lower()
        CTX.announce(f"{pkg} is installed, {config.USER_TITLE}." if ok else
                     f"The {pkg} installation didn't succeed, {config.USER_TITLE}.")
        if ok and CTX.apps:
            threading.Thread(target=CTX.apps.refresh, daemon=True).start()
    threading.Thread(target=run, daemon=True).start()
    return f"Installing {pkg} ({pid}) in the background; I'll announce when it's done."


# =============================================================================
# Protocols (routines) and personas
# =============================================================================
ROUTINES_FILE = DATA / "routines.json"
DEFAULT_ROUTINES = {
    "work": {"steps": ["open Visual Studio Code", "open Claude", "play lofi beats on Spotify", "volume 30"],
             "say": "Work protocol engaged. Let's build something remarkable."},
    "gaming": {"steps": ["open Steam", "open Discord", "volume 60"],
               "say": "Gaming protocol engaged. Good hunting."},
    "focus": {"steps": ["minimize everything", "play deep focus on Spotify", "volume 25"],
              "say": "Focus protocol engaged. I'll keep things quiet."},
    "night": {"steps": ["dark mode", "pause the music", "volume 15"], "say": "Night protocol engaged. Rest well."},
}


def load_routines() -> dict:
    try:
        return json.loads(ROUTINES_FILE.read_text(encoding="utf-8"))
    except Exception:
        ROUTINES_FILE.write_text(json.dumps(DEFAULT_ROUTINES, indent=1), encoding="utf-8")
        return dict(DEFAULT_ROUTINES)


def find_routine(name: str):
    rs = load_routines()
    key = re.sub(r"\b(protocol|routine|mode|the|my)\b", "", name.lower()).strip()
    if key in rs:
        return key, rs[key]
    close = [k for k in rs if k in key or key in k]
    return (close[0], rs[close[0]]) if close else (None, None)


@tool("Protocols (saved routines): run one by name, save a new one as a list of plain-English commands, "
      "list them, or delete one. E.g. save 'study' = ['close Discord', 'play piano focus on Spotify', 'volume 20'].",
      {"action": S("Action", ["run", "save", "list", "delete"]), "name": S("Protocol name"),
       "steps": {"type": "array", "items": {"type": "string"}, "description": "For save: commands in order"},
       "say": S("For save: what to say when it runs")}, ["action"],
      confirm=lambda a: f"Delete the {a.get('name')} protocol?" if a.get("action") == "delete" else None)
def routine(action: str, name: str = "", steps: list = None, say: str = ""):
    rs = load_routines()
    if action == "list":
        return "Protocols: " + "; ".join(f"{k} ({', '.join(v['steps'])})" for k, v in rs.items())
    if action == "save":
        if not name or not steps:
            return "Need a name and the steps."
        key = re.sub(r"\b(protocol|routine|mode)\b", "", name.lower()).strip()
        rs[key] = {"steps": [str(s) for s in steps], "say": say or f"{key.title()} protocol engaged."}
        ROUTINES_FILE.write_text(json.dumps(rs, indent=1), encoding="utf-8")
        return f"Saved the {key} protocol with {len(steps)} steps."
    key, r = find_routine(name)
    if not r:
        return f"No protocol called {name}. Existing: {', '.join(rs)}."
    if action == "delete":
        rs.pop(key, None)
        ROUTINES_FILE.write_text(json.dumps(rs, indent=1), encoding="utf-8")
        return f"Deleted the {key} protocol."
    results = CTX.run_steps(r["steps"])
    failed = [s for s, res in zip(r["steps"], results) if res.lower().startswith(("error", "couldn", "no "))]
    return r.get("say", f"{key.title()} protocol engaged.") + (f" Some steps had trouble: {'; '.join(failed)}."
                                                              if failed else "")


@tool("Switch persona: 'jarvis' (British gentleman, calls you sir) or 'friday' (warm and quick, calls you boss).",
      {"name": S("Persona", ["jarvis", "friday"])}, ["name"])
def switch_persona(name: str):
    from . import persona
    return persona.apply(name)
